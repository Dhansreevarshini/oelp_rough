from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ============================================================
# Place this script in manifesto_analysis/ (next to the other scripts).
# It reads the gemini.csv files + the original chunk .txt files and builds
# one Excel file per dataset, with the chunk text next to the model's labels.
# It does NOT call the Gemini API, so it costs nothing and can be re-run anytime.
# ============================================================

ROOT = Path(__file__).resolve().parent
PARTIES = ["AIADMK", "DMK", "TVK"]

DATASETS = {
    "speech": {
        "chunks_dir": ROOT / "chunks_speech",
        "results_dir": ROOT / "output" / "speech_results",
        "csv_name": "gemini.csv",
        "xlsx_path": ROOT / "output" / "speech_results" / "speech_classification_review.xlsx",
    },
    "manifesto": {
        "chunks_dir": ROOT / "chunks",
        "results_dir": ROOT / "output" / "manifesto_results",
        "csv_name": "gemini.csv",
        "xlsx_path": ROOT / "output" / "manifesto_results" / "manifesto_classification_review.xlsx",
    },
}

HEADER_SEPARATOR = "-" * 40
TEXT_FONT = "Nirmala UI"     # supports Tamil on Windows; Excel falls back if missing
BASE_FONT = "Arial"

COLUMNS = [
    "Party", "Chunk", "Primary_Category", "Secondary_Category",
    "Chunk_Text", "Status", "Error", "Verified_Correct", "Reviewer_Notes",
]
WIDTHS = {
    "Party": 10, "Chunk": 42, "Primary_Category": 30, "Secondary_Category": 30,
    "Chunk_Text": 90, "Status": 10, "Error": 25, "Verified_Correct": 16, "Reviewer_Notes": 35,
}


def strip_header(raw_text):
    """Remove the 'Party folder / Source / Chunk / Words' header if present."""
    lines = raw_text.splitlines()
    for i, line in enumerate(lines):
        if line.strip() == HEADER_SEPARATOR:
            return "\n".join(lines[i + 1:]).strip()
    return raw_text.strip()


def read_chunk_text(chunks_dir, party, chunk_name):
    path = chunks_dir / party / chunk_name
    if not path.exists():
        return "[chunk file not found]"
    return strip_header(path.read_text(encoding="utf-8", errors="replace"))


def build_party_frame(cfg, party):
    csv_path = cfg["results_dir"] / party / cfg["csv_name"]
    if not csv_path.exists():
        print(f"  {party}: no CSV found ({csv_path.relative_to(ROOT)}), skipped")
        return None

    df = pd.read_csv(csv_path, encoding="utf-8-sig", keep_default_na=False)
    if df.empty:
        print(f"  {party}: CSV is empty, skipped")
        return None

    df["Chunk_Text"] = [read_chunk_text(cfg["chunks_dir"], party, c) for c in df["Chunk"]]
    df["Verified_Correct"] = ""
    df["Reviewer_Notes"] = ""
    for col in COLUMNS:
        if col not in df.columns:
            df[col] = ""
    return df[COLUMNS]


def style_sheet(ws, n_rows):
    header_fill = PatternFill("solid", start_color="1F3864")
    for cell in ws[1]:
        cell.font = Font(name=BASE_FONT, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = WIDTHS[col]

    text_col = COLUMNS.index("Chunk_Text") + 1
    for row in ws.iter_rows(min_row=2, max_row=n_rows + 1):
        for cell in row:
            is_text = cell.column == text_col
            cell.font = Font(name=TEXT_FONT if is_text else BASE_FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Dropdown Yes / No in the Verified_Correct column
    verify_col = get_column_letter(COLUMNS.index("Verified_Correct") + 1)
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{verify_col}2:{verify_col}{n_rows + 1}")

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{n_rows + 1}"


def build_dataset(name, cfg):
    print(f"\n{name.upper()}")
    frames = {}
    for party in PARTIES:
        df = build_party_frame(cfg, party)
        if df is not None:
            frames[party] = df
            print(f"  {party}: {len(df)} chunks")

    if not frames:
        print(f"  No results found for {name}, no Excel file created.")
        return

    cfg["xlsx_path"].parent.mkdir(parents=True, exist_ok=True)
    all_df = pd.concat(frames.values(), ignore_index=True)

    with pd.ExcelWriter(cfg["xlsx_path"], engine="openpyxl") as writer:
        all_df.to_excel(writer, sheet_name="All", index=False)
        for party, df in frames.items():
            df.to_excel(writer, sheet_name=party, index=False)

        for sheet_name, df in [("All", all_df)] + list(frames.items()):
            style_sheet(writer.sheets[sheet_name], len(df))

    print(f"  Saved: {cfg['xlsx_path'].relative_to(ROOT)}")


def main():
    for name, cfg in DATASETS.items():
        build_dataset(name, cfg)


if __name__ == "__main__":
    main()